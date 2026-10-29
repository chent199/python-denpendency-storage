"""Run inside a fresh target Python venv after offline installation."""
import asyncio
import io
import os
import tempfile
import threading
import socket
import time
import pathlib
import aiofiles
import httpx
import requests
import uvicorn
from openai import OpenAI
from langchain_openai import ChatOpenAI
from pydantic import BaseModel, ValidationError
from openpyxl import Workbook, load_workbook
from docx import Document
from dotenv import load_dotenv
from loguru import logger
from fastapi import FastAPI, UploadFile, File
from fastapi.testclient import TestClient
from deepagents import create_deep_agent
from langchain_core.language_models.fake_chat_models import FakeListChatModel

app = FastAPI()


@app.get('/probe')
def probe():
    return {'offline': True}


assert TestClient(app).get('/probe').json() == {'offline': True}


@app.post('/upload')
async def upload(file: UploadFile = File()):
    return {'text': (await file.read()).decode()}


assert TestClient(app).post('/upload', files={'file': ('test.txt', b'upload-ok')}).json() == {'text': 'upload-ok'}


async def file_test():
    with tempfile.TemporaryDirectory() as folder:
        path = pathlib.Path(folder) / 'test.txt'
        async with aiofiles.open(path, 'w') as stream:
            await stream.write('aiofiles-ok')
        async with aiofiles.open(path) as stream:
            assert await stream.read() == 'aiofiles-ok'


asyncio.run(file_test())
with tempfile.TemporaryDirectory() as folder:
    xlsx = pathlib.Path(folder) / 'test.xlsx'
    book = Workbook()
    book.active['A1'] = 'excel-ok'
    book.save(xlsx)
    loaded = load_workbook(xlsx)
    assert loaded.active['A1'].value == 'excel-ok'
    loaded.close()
    docx = pathlib.Path(folder) / 'test.docx'
    document = Document()
    document.add_paragraph('word-ok')
    document.save(docx)
    assert Document(docx).paragraphs[0].text == 'word-ok'
load_dotenv(stream=io.StringIO('OFFLINE_SMOKE_VALUE=dotenv-ok\n'))
assert os.environ['OFFLINE_SMOKE_VALUE'] == 'dotenv-ok'
sink = io.StringIO()
handle = logger.add(sink, format='{message}')
logger.info('loguru-ok')
logger.remove(handle)
assert 'loguru-ok' in sink.getvalue()

# A real loopback HTTP request validates uvicorn and both HTTP clients without Internet.
sock = socket.socket()
sock.bind(('127.0.0.1', 0))
port = sock.getsockname()[1]
server = uvicorn.Server(uvicorn.Config(app, log_level='error', lifespan='off'))
thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
thread.start()
try:
    for attempt in range(100):
        if server.started:
            break
        time.sleep(0.05)
    else:
        raise RuntimeError('uvicorn failed to start')
    url = f'http://127.0.0.1:{port}/probe'
    assert requests.get(url, timeout=5).json() == {'offline': True}
    assert httpx.get(url, timeout=5).json() == {'offline': True}
    # Local OpenAI-compatible endpoint exercises SDK and LangChain serialization.
    @app.post('/v1/chat/completions')
    def completion():
        return {'id': 'offline-test', 'object': 'chat.completion', 'created': 0,
                'model': 'offline-model', 'choices': [{'index': 0,
                'message': {'role': 'assistant', 'content': 'sdk-ok'}, 'finish_reason': 'stop'}],
                'usage': {'prompt_tokens': 1, 'completion_tokens': 1, 'total_tokens': 2}}

    base = f'http://127.0.0.1:{port}/v1'
    with OpenAI(api_key='offline-test-placeholder', base_url=base, max_retries=0) as client:
        answer = client.chat.completions.create(model='offline-model', messages=[{'role': 'user', 'content': 'hello'}])
        assert answer.choices[0].message.content == 'sdk-ok'
    with httpx.Client() as client:
        model = ChatOpenAI(model='offline-model', api_key='offline-test-placeholder', base_url=base,
                           max_retries=0, http_client=client, tiktoken_model_name='gpt-4o')
        assert model.invoke('hello').content == 'sdk-ok'
        assert model.get_num_tokens('离线 token counting test') > 0
finally:
    server.should_exit = True
    thread.join(timeout=10)
    sock.close()
assert not thread.is_alive()


class Payload(BaseModel):
    count: int


assert Payload.model_validate({'count': '3'}).model_dump() == {'count': 3}
try:
    Payload.model_validate({'count': 'invalid'})
except ValidationError:
    pass
else:
    raise AssertionError('Pydantic must reject invalid data')


class OfflineChat(FakeListChatModel):
    def bind_tools(self, tools, **kwargs):
        return self


agent = create_deep_agent(model=OfflineChat(responses=['offline-ok']))
result = agent.invoke({'messages': [{'role': 'user', 'content': 'hello'}]})
assert result['messages'][-1].content == 'offline-ok'
print('FastAPI HTTP request and deepagents graph invocation: PASS (local fake model, no network)')
print('aiofiles, httpx, loguru, python-dotenv, python-multipart, requests, uvicorn: PASS')
print('openai and langchain-openai local HTTP completions, pydantic validation: PASS')
print('openpyxl XLSX and python-docx DOCX round trips: PASS')

assert os.environ.get('TIKTOKEN_CACHE_DIR'), 'Set TIKTOKEN_CACHE_DIR to the delivered vocabulary directory'
import tiktoken
for name in tiktoken.list_encoding_names():
    encoding = tiktoken.get_encoding(name)
    assert encoding.decode(encoding.encode('离线 tokenizer test')) == '离线 tokenizer test'
print('All tiktoken encodings and LangChain token counting offline: PASS')

import pytest
result = pytest.main(['-q', '-p', 'no:cacheprovider', str(pathlib.Path(__file__).with_name('test_additions.py'))])
assert result == 0, 'Additional package functional checks failed'
